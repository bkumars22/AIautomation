"""
Step 6 verification. The explicit, named requirement: "downloads must
contain real, matching data -- verify the downloaded .docx manual cases
and .zip automated code actually correspond to the same session, not
stale or mismatched content." These tests build a bundle from known
input and then READ THE BUNDLE BACK (via python-docx / zipfile, not by
trusting the builder), asserting the exact content round-trips.
"""
from __future__ import annotations

import io
import zipfile

from docx import Document as DocxDocument

from document_upload.downloads import build_automated_code_zip, build_manual_cases_docx

_TESTS = [
    {
        "scenario_name": "Apply a valid promo code",
        "manual_test_case": "Test Case: Apply a valid promo code\nSteps:\n1. Go to /checkout.",
        "automated_code": "def test_apply_promo(): pass",
        "framework": "playwright",
    },
    {
        "scenario_name": "Reject an invalid promo code",
        "manual_test_case": "Test Case: Reject an invalid promo code\nSteps:\n1. Go to /checkout.",
        "automated_code": "describe('reject promo', () => {});",
        "framework": "cypress",
    },
]


def test_docx_contains_every_scenario_heading_and_body():
    docx_bytes = build_manual_cases_docx(_TESTS)
    doc = DocxDocument(io.BytesIO(docx_bytes))
    full_text = "\n".join(p.text for p in doc.paragraphs)

    for test in _TESTS:
        assert test["scenario_name"] in full_text
        assert test["manual_test_case"] in full_text


def test_docx_does_not_mix_up_scenario_order():
    docx_bytes = build_manual_cases_docx(_TESTS)
    doc = DocxDocument(io.BytesIO(docx_bytes))
    full_text = "\n".join(p.text for p in doc.paragraphs)

    first_pos = full_text.index(_TESTS[0]["scenario_name"])
    second_pos = full_text.index(_TESTS[1]["scenario_name"])
    assert first_pos < second_pos


def test_zip_contains_one_file_per_scenario_with_correct_extension():
    zip_bytes = build_automated_code_zip(_TESTS)
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        names = zf.namelist()
        assert any(n.endswith(".py") for n in names)  # playwright
        assert any(n.endswith(".js") for n in names)  # cypress
        assert len(names) == 2


def test_zip_file_content_matches_the_right_scenario_not_swapped():
    zip_bytes = build_automated_code_zip(_TESTS)
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        for name in zf.namelist():
            content = zf.read(name).decode("utf-8")
            if name.endswith(".py"):
                assert content == _TESTS[0]["automated_code"]
            elif name.endswith(".js"):
                assert content == _TESTS[1]["automated_code"]


def test_zip_handles_unknown_framework_extension_gracefully():
    tests = [{"scenario_name": "X", "manual_test_case": "m", "automated_code": "c", "framework": "unknown-tool"}]
    zip_bytes = build_automated_code_zip(tests)
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        assert zf.namelist() == ["X.txt"]


def test_filename_sanitizes_unsafe_characters():
    tests = [{"scenario_name": "User can log in / reset: password?", "manual_test_case": "m", "automated_code": "c", "framework": "playwright"}]
    zip_bytes = build_automated_code_zip(tests)
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        name = zf.namelist()[0]
        assert "/" not in name.replace(".py", "")
        assert ":" not in name
        assert "?" not in name
