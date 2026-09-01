"""
Exports a list of UnifiedTestResult to formats a real QA team actually
shares: CSV (universal, opens in any spreadsheet tool) and .xlsx (native
Excel, with a status column colored for at-a-glance scanning — the kind
of thing a manager expects when someone says "here's the test report").
"""
from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from reporting.unified_report import UnifiedTestResult

_COLUMNS = ["test_name", "tool_used", "status", "duration_ms", "error_message"]

_STATUS_FILL = {
    "passed": PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid"),
    "failed": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
    "skipped": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
}


def export_to_csv(results: list[UnifiedTestResult]) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(_COLUMNS)
    for r in results:
        writer.writerow([r.test_name, r.tool_used, r.status, r.duration_ms, r.error_message or ""])
    return buffer.getvalue()


def export_to_excel(results: list[UnifiedTestResult]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Results"

    header = ["Test Name", "Tool", "Status", "Duration (ms)", "Error Message"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for r in results:
        ws.append([r.test_name, r.tool_used, r.status, r.duration_ms, r.error_message or ""])
        row_idx = ws.max_row
        fill = _STATUS_FILL.get(r.status)
        if fill:
            ws.cell(row=row_idx, column=3).fill = fill

    passed = sum(1 for r in results if r.status == "passed")
    ws.append([])
    ws.append(["Summary", f"{passed}/{len(results)} passed"])
    ws["A" + str(ws.max_row)].font = Font(bold=True)

    for col_idx, width in enumerate([45, 12, 10, 14, 60], start=1):
        ws.column_dimensions[chr(64 + col_idx)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
