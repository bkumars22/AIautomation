"""
Step 6 of the document-upload feature: bundle every generated test for
one session into a single downloadable file. Pure functions (dict list
in, bytes out) -- no FastAPI, no database -- so downloads/*.py can be
tested directly against real generated content without touching a
session or the audit DB.
"""
from __future__ import annotations

import io
import zipfile
from typing import Any

from docx import Document as DocxDocument

_EXTENSION_BY_FRAMEWORK = {
    "playwright": "py",
    "selenium": "py",
    "cypress": "js",
    "testng-selenium": "java",
    "cucumber": "feature",
}


def _safe_filename(name: str) -> str:
    return "".join(c if c.isalnum() or c in " _-" else "_" for c in name).strip().replace(" ", "_")


def build_manual_cases_docx(tests: list[dict[str, Any]]) -> bytes:
    doc = DocxDocument()
    doc.add_heading("Manual Test Cases", 0)
    for test in tests:
        doc.add_heading(test["scenario_name"], level=1)
        doc.add_paragraph(test["manual_test_case"])
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_automated_code_zip(tests: list[dict[str, Any]]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as zf:
        for test in tests:
            ext = _EXTENSION_BY_FRAMEWORK.get(test["framework"], "txt")
            zf.writestr(f"{_safe_filename(test['scenario_name'])}.{ext}", test["automated_code"])
    return buffer.getvalue()
