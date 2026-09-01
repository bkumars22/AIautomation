"""One-off script: generates tests/fixtures/documents/business_requirements_profile.docx.
Run once (`python scripts/generate_docx_fixture.py` from the repo root) to
(re)produce the fixture; not imported at runtime."""
from pathlib import Path

from docx import Document

OUTPUT = Path(__file__).parent.parent / "tests" / "fixtures" / "documents" / "business_requirements_profile.docx"

doc = Document()
doc.add_heading("Business Requirements: User Profile Management", level=0)

doc.add_paragraph(
    "This document describes the user profile management capabilities being "
    "added to the customer portal in Q2. It is intended for the QA team to "
    "derive test coverage from, alongside the engineering implementation."
)

doc.add_heading("Background", level=1)
doc.add_paragraph(
    "Customers have long requested the ability to manage their own contact "
    "details without contacting support. Currently, only support staff can "
    "update a customer's email or phone number via the internal admin tool."
)

doc.add_heading("Requirement: Self-service email update", level=1)
doc.add_paragraph(
    "A logged-in customer should be able to open their account settings page "
    "and change the email address on file. Because email changes affect "
    "login, the system must first send a confirmation link to the NEW "
    "address; the change should not take effect until that link is clicked. "
    "If the customer never confirms, the old email remains active and no "
    "error is shown to anyone else — this is expected, not a bug."
)

doc.add_heading("Requirement: Password strength feedback", level=1)
doc.add_paragraph(
    "When a customer sets a new password from the account settings page, "
    "the form should show a live strength indicator (weak/medium/strong) "
    "as they type, and should refuse to submit a password rated weak. The "
    "same rules apply whether the customer is setting a password for the "
    "first time or changing an existing one."
)

doc.add_heading("Requirement: Profile photo upload size limit", level=1)
doc.add_paragraph(
    "Customers may upload a profile photo. Files larger than 5 MB must be "
    "rejected with a clear message before any upload is attempted, rather "
    "than failing partway through. Accepted formats are JPG and PNG only."
)

doc.add_heading("Non-functional notes", level=1)
doc.add_paragraph(
    "All of the above must work on both desktop and mobile web layouts. "
    "This section does not describe a testable user action on its own."
)

doc.save(str(OUTPUT))
print("wrote", OUTPUT)
